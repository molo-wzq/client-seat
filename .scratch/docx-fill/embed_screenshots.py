from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.shared import Inches, Pt
from docx.oxml.ns import qn


SOURCE = Path(r"D:\话术对练\作品情况填写版.docx")
OUTPUT = Path(r"D:\话术对练\作品情况填写版_含截图.docx")
BASE = Path(r"D:\话术对练\.scratch\docx-fill\作品情况填写版-无截图.docx")

FIGURES = [
    (
        Path(r"D:\话术对练\docs\ui-setup.png"),
        "图1 布置幕 选择生客入座并查看本局策略卡",
        "Client Seat布置幕，展示生客画像、客户席、策略卡和产品卡",
    ),
    (
        Path(r"D:\话术对练\docs\ui-table.png"),
        "图2 对局幕 AI根据客户回应自动使用已发布策略卡",
        "Client Seat对局幕，展示客户信息、模拟通话和本轮使用的策略卡",
    ),
    (
        Path(r"D:\话术对练\docs\ui-postgame.png"),
        "图3 结算页 按轮次查看策略路径和来源片段",
        "Client Seat结算页，展示沟通结果、结束原因和策略路径",
    ),
    (
        Path(r"D:\话术对练\docs\ui-cards.png"),
        "图4 策略卡列表 已发布卡才会进入对局",
        "Client Seat策略卡列表，展示四张已发布策略卡",
    ),
]


def insert_paragraph_after(anchor: Paragraph, style: str | None = None) -> Paragraph:
    element = OxmlElement("w:p")
    anchor._p.addnext(element)
    paragraph = Paragraph(element, anchor._parent)
    if style:
        paragraph.style = style
    return paragraph


def set_chinese_font(run, name: str, size: float) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)


def set_alt_text(run, description: str) -> None:
    for properties in run._element.xpath(".//pic:cNvPr | .//wp:docPr"):
        properties.set("descr", description)
        properties.set("title", description)


def main() -> None:
    if not BASE.exists():
        shutil.copy2(SOURCE, BASE)

    document = Document(BASE)
    screenshot_line = next(
        (
            paragraph
            for paragraph in document.paragraphs
            if paragraph.text.startswith("3.主要截图：")
        ),
        None,
    )
    if screenshot_line is None:
        raise RuntimeError("Screenshot placeholder paragraph not found")

    screenshot_line.clear()
    title_run = screenshot_line.add_run("3.主要功能截图如下：")
    set_chinese_font(title_run, "仿宋_GB2312", 16)

    following_materials = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text.startswith("4.配套材料：")
    )
    following_materials.paragraph_format.page_break_before = True

    anchor = screenshot_line
    for index, (image_path, caption, alt_text) in enumerate(FIGURES):
        image_paragraph = insert_paragraph_after(anchor, "Normal")
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.paragraph_format.space_before = Pt(4)
        image_paragraph.paragraph_format.space_after = Pt(2)
        image_paragraph.paragraph_format.keep_with_next = True
        if index in (0, 2):
            image_paragraph.paragraph_format.page_break_before = True
        image_run = image_paragraph.add_run()
        image_run.add_picture(str(image_path), width=Inches(5.25))
        set_alt_text(image_run, alt_text)

        caption_paragraph = insert_paragraph_after(image_paragraph, "Normal")
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_paragraph.paragraph_format.space_before = Pt(0)
        caption_paragraph.paragraph_format.space_after = Pt(8)
        caption_paragraph.paragraph_format.keep_together = True
        caption_run = caption_paragraph.add_run(caption)
        set_chinese_font(caption_run, "仿宋_GB2312", 11)
        anchor = caption_paragraph

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
